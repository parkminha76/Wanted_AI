"""파일 -> raw_text + 서식 정보(spans). 이미지/텍스트 라우팅도 여기서 한다.

B-1(파일 형식)과 B-2(탐지) 사이의 유일한 약속이 이 파일의 출력 형식이다.
`ParsedDoc` / `TextSpan`만 지키면 두 사람이 서로 안 보고 작업할 수 있다.

이 파일이 지키는 것
-------------------
1. **원문을 고치지 않는다.** 줄바꿈 정규화도, BOM 제거도 하지 않는다. Finding의
   offset이 raw_text 기준이고 마스킹이 그 offset을 되짚기 때문에, 여기서 한 글자라도
   손대면 화면 하이라이트와 리댁션이 전부 어긋난다.
2. **span은 잘게 쪼갠다.** txt는 줄 단위, docx는 run 단위, xlsx는 셀 단위, PDF는
   texttrace span 단위다. hidden.py가 "제로폭 밀도"를 잴 때 분모가 문서 전체가 되면
   5만 자짜리 계약서에 200개를 심어도 0.4%로 통과한다 (README 함정 1).
3. **모르는 서식 값은 "탐지가 안 되는 쪽"으로 채운다.** 색은 검정, 크기는 11pt,
   배경은 흰색이다. 모른다는 이유로 오탐이 나면 안 된다. 다만 배경색만은
   "모르는 값"과 "진짜 흰색"을 구분해야 해서 `bg_known`을 따로 둔다.

README 계약에 없는 필드를 몇 개 더 붙였다(추가는 자유, 기존 필드는 그대로다).
`bg_known` · `where` · `ParsedDoc.path` · `ParsedDoc.file_type` — 아래 각 정의 참고.
"""

from __future__ import annotations

import colorsys
import os
import re
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass, field


# 서식 값을 알아낼 수 없을 때 쓰는 기본값.
#
# 전부 "hidden.py의 검사를 통과하는 쪽"으로 잡혀 있다. 폰트 크기를 0.0으로 두면
# hidden.py의 "2pt 미만" 규칙에 전부 걸려서 멀쩡한 문서가 빨간불이 된다.
# 모르면 놓치는 게 낫지, 모른다고 신고하면 안 된다.
DEFAULT_FONT_SIZE = 11.0
DEFAULT_COLOR = "#000000"
DEFAULT_BG_COLOR = "#ffffff"


@dataclass
class TextSpan:
    """서식 정보가 붙은 텍스트 조각 하나. hidden.py가 이것만 본다."""

    text: str
    start: int                  # raw_text 기준 offset
    end: int
    page: int                   # PDF는 페이지, XLSX는 시트 번호, 그 외는 1
    font_size: float = DEFAULT_FONT_SIZE   # pt
    color: str = DEFAULT_COLOR             # "#ffffff"
    bg_color: str = DEFAULT_BG_COLOR       # 그 위치의 실제 배경 채움색
    render_mode: int = 0        # PDF 전용. 3 = 화면에 안 그려짐
    opacity: float = 1.0        # PDF 전용. 0.0 = 완전 투명
    hidden_attr: bool = False   # docx의 vanish, xlsx의 숨긴 행/열·시트 등
    bbox: tuple | None = None   # 마스킹 좌표 (PDF만 채워진다)

    # 글자별 가로 좌표. 글자 수 + 1개다 — 각 글자의 왼쪽 경계에 마지막 글자의
    # 오른쪽 경계를 하나 더 붙인 것. i번째 글자가 차지하는 x 구간은
    # char_x[i] ~ char_x[i+1]이다. PDF 가로쓰기에서만 채워진다.
    #
    # 여기 있는 이유: 마스킹은 "몇 번째 글자"가 아니라 "페이지 위 좌표"로 지운다.
    # span 사각형 하나만 남기면 그 안의 특정 구간 좌표를 글자 수로 비례 배분해
    # 짐작할 수밖에 없는데, 한글(약 10.3pt)과 ASCII(약 6.7pt)는 폭이 1.5배 차이라
    # 섞인 줄에서 글자 한두 개분씩 밀린다. 덜 덮이면 개인정보가 그대로 남고,
    # 넓히면 옆 글자까지 지워진다. get_texttrace()가 글자마다 주는 좌표를 그냥
    # 들고 있으면 짐작할 필요가 없어진다.
    #
    # 세로 좌표를 안 싣는 이유: 같은 span 안의 글자는 y 구간을 공유한다.
    # bbox[1], bbox[3]을 그대로 쓰면 되므로 숫자를 4분의 1로 줄인다.
    char_x: list[float] | None = None

    # --- 아래 둘은 README 계약에 없는 추가 필드다 ---

    # bg_color가 진짜 그 자리의 채움색인지, 알아내지 못해서 흰색으로 둔 것인지.
    # False면 "모른다"는 뜻이다. 이걸 구분하지 못하면 hidden.py의 "글자색 ≈ 배경색"
    # 검사가 흰 글씨를 무조건 신고하게 되고, 흰 배경이 아닌 문서에서 오탐이 난다.
    bg_known: bool = False

    # 이 조각이 문서의 어디서 나왔는지. "body" | "header" | "footer" | "textbox"
    # | "comment" | "footnote" | "deleted" | "cell" | "sheet_hidden" 등.
    #
    # 머리말·꼬리말·주석을 여기서 hidden_attr=True로 찍지 않는 이유: 모든 문서에
    # 머리말이 있다. 서식상 숨겨진 것이 아니라 "눈에 잘 안 띄는 자리"일 뿐이라,
    # 판정은 hidden.py가 다른 근거와 함께 내려야 한다. 여기서는 자리만 알려준다.
    where: str = "body"

    # hidden_attr이 True가 된 이유. 여러 개면 "+"로 잇는다.
    #   "vanish" | "web_hidden" | "deleted" | "row_hidden" | "col_hidden"
    #   | "sheet_hidden" | "sheet_very_hidden" | "blank_format"
    #
    # 여기 있는 이유: 같은 hidden_attr=True라도 무게가 전혀 다르다. 숨긴 행·열은
    # 보조 계산용으로 정상 문서에 흔하지만, `;;;` 서식과 veryHidden 시트는 실수로
    # 만들어지지 않는다. 하나로 뭉뚱그리면 hidden.py가 확신도를 나눌 근거를 잃는다.
    hidden_reason: str = ""


@dataclass
class ParsedDoc:
    """파일 하나를 푼 결과. scan.py가 raw_text와 spans를 가져간다."""

    kind: str                   # "text" | "image" — 어느 파이프라인으로 보낼지
    raw_text: str               # 문서 전체 텍스트. 모든 offset의 기준이다.
    spans: list[TextSpan] = field(default_factory=list)
    page_map: list[int] = field(default_factory=list)   # 문자 1개당 페이지 번호 1개

    # --- 추가 필드 ---
    path: str = ""              # 원본 경로. 이미지 파이프라인(ml/의 CNN)이 이걸 받는다.
    file_type: str = ""         # ScanResult.file_type에 그대로 들어간다.


class ParseError(Exception):
    """파일을 열거나 읽을 수 없다. scan.py가 잡아서 ScanResult.error로 옮긴다.

    메시지에 파일 내용을 넣지 않는다 (로그에 원문을 남기지 않는다 — 팀 규칙 2).
    """


# 확장자 -> file_type. 여기 없는 확장자는 ParseError다.
TEXT_EXTENSIONS: dict[str, str] = {
    ".txt": "txt",
    ".md": "md",
    ".csv": "csv",
    ".log": "txt",
}
IMAGE_EXTENSIONS: dict[str, str] = {
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".bmp": "image",
    ".gif": "image",
    ".webp": "image",
    ".tif": "image",
    ".tiff": "image",
}

# 구버전 바이너리 형식. 라이브러리가 아예 못 읽으므로 안내만 한다.
_LEGACY_EXTENSIONS = {".doc": "docx", ".xls": "xlsx", ".ppt": "pptx"}

# 텍스트 레이어가 이만큼도 안 나오는 PDF는 스캔본으로 보고 이미지 파이프라인에 넘긴다.
_PDF_SCANNED_TEXT_THRESHOLD = 20


def load(path: str) -> ParsedDoc:
    """파일 하나를 ParsedDoc으로 푼다. scan.py의 유일한 진입점이다."""
    if not os.path.isfile(path):
        raise ParseError("파일이 없다")

    ext = os.path.splitext(path)[1].lower()

    if ext in TEXT_EXTENSIONS:
        return _load_text_file(path, TEXT_EXTENSIONS[ext])
    if ext in IMAGE_EXTENSIONS:
        return _load_image(path)
    if ext == ".pdf":
        return _load_pdf(path)
    if ext == ".docx":
        return _load_docx(path)
    if ext in (".xlsx", ".xlsm"):
        return _load_xlsx(path)
    if ext in _LEGACY_EXTENSIONS:
        raise ParseError(
            f"구버전 형식({ext})은 지원하지 않는다. {_LEGACY_EXTENSIONS[ext]}로 변환해 달라"
        )
    raise ParseError(f"지원하지 않는 형식({ext or '확장자 없음'})")


# ---------------------------------------------------------------------------
# offset 누적기 — 이 파일의 핵심이자 유일한 어려움
# ---------------------------------------------------------------------------


class _Builder:
    """raw_text를 이어붙이면서 각 조각의 offset과 페이지를 같이 기록한다.

    여기가 어긋나면 화면 하이라이트가 엉뚱한 글자에 칠해지고 리댁션이 멀쩡한 글자를
    지운다. 그래서 raw_text에 들어가는 모든 문자는 반드시 이 클래스를 거친다.
    """

    def __init__(self) -> None:
        self._parts: list[str] = []
        self._pages: list[int] = []
        self.spans: list[TextSpan] = []
        self.cursor = 0

    def add_span(self, text: str, page: int, **fmt) -> TextSpan | None:
        """텍스트 조각 하나를 raw_text에 붙이고 span으로 기록한다."""
        if not text:
            return None
        span = TextSpan(
            text=text, start=self.cursor, end=self.cursor + len(text), page=page, **fmt
        )
        self.spans.append(span)
        self._append(text, page)
        return span

    def add_gap(self, text: str, page: int) -> None:
        """구분자(줄바꿈·공백·탭)를 붙인다. 문서에 있던 글자가 아니므로 span은 만들지 않는다.

        span으로 만들면 hidden.py가 공백 한 칸짜리 조각을 서식 검사하게 되고,
        마스킹이 구분자 위치를 값의 일부로 착각한다.
        """
        if not text:
            return
        self._append(text, page)

    def newline(self, page: int) -> None:
        """줄바꿈. 단, 맨 앞에는 넣지 않는다 (offset이 1씩 밀린다)."""
        if self.cursor:
            self.add_gap("\n", page)

    def _append(self, text: str, page: int) -> None:
        self._parts.append(text)
        self._pages.extend([page] * len(text))
        self.cursor += len(text)

    def build(self, kind: str, path: str, file_type: str) -> ParsedDoc:
        raw_text = "".join(self._parts)
        assert len(self._pages) == len(raw_text)   # offset 회계가 맞는지 확인
        return ParsedDoc(
            kind=kind,
            raw_text=raw_text,
            spans=self.spans,
            page_map=self._pages,
            path=path,
            file_type=file_type,
        )


# ---------------------------------------------------------------------------
# 색 유틸
# ---------------------------------------------------------------------------


def _rgb_floats_to_hex(r: float, g: float, b: float) -> str:
    def channel(value: float) -> int:
        return max(0, min(255, round(value * 255)))

    return f"#{channel(r):02x}{channel(g):02x}{channel(b):02x}"


def _pdf_color_to_hex(color) -> str:
    """PyMuPDF의 색 값(0~1 실수 튜플)을 #rrggbb로. 성분 개수로 색공간을 판단한다."""
    if color is None:
        return DEFAULT_COLOR
    if isinstance(color, (int, float)):
        components = [float(color)]
    else:
        try:
            components = [float(c) for c in color]
        except (TypeError, ValueError):
            return DEFAULT_COLOR

    if len(components) == 1:                      # gray
        value = components[0]
        return _rgb_floats_to_hex(value, value, value)
    if len(components) == 3:                      # rgb
        return _rgb_floats_to_hex(*components)
    if len(components) == 4:                      # cmyk
        c, m, y, k = components
        return _rgb_floats_to_hex((1 - c) * (1 - k), (1 - m) * (1 - k), (1 - y) * (1 - k))
    return DEFAULT_COLOR


def _hex_from_argb(value) -> str | None:
    """엑셀/워드의 ARGB("FFFFFFFF") 또는 RGB("FFFFFF") 문자열을 #rrggbb로."""
    if not isinstance(value, str):
        return None
    digits = value.strip().lstrip("#")
    if len(digits) == 8:
        digits = digits[2:]
    if not re.fullmatch(r"[0-9A-Fa-f]{6}", digits):
        return None
    return "#" + digits.lower()


def _apply_tint(hex_color: str, tint: float | None) -> str:
    """OOXML의 tint(명도 보정). 엑셀의 "흰색, 배경 1, 25% 더 어둡게"가 이것이다."""
    if not tint:
        return hex_color
    r = int(hex_color[1:3], 16) / 255
    g = int(hex_color[3:5], 16) / 255
    b = int(hex_color[5:7], 16) / 255
    hue, lightness, saturation = colorsys.rgb_to_hls(r, g, b)
    lightness = lightness * (1 + tint) if tint < 0 else lightness * (1 - tint) + tint
    return _rgb_floats_to_hex(
        *colorsys.hls_to_rgb(hue, max(0.0, min(1.0, lightness)), saturation)
    )


# ---------------------------------------------------------------------------
# TXT / MD / CSV
# ---------------------------------------------------------------------------


def _decode(data: bytes) -> str:
    """UTF-8 -> cp949 순서로 시도한다. 둘 다 실패하면 깨진 글자만 대체한다."""
    for encoding in ("utf-8", "cp949"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _load_text_file(path: str, file_type: str) -> ParsedDoc:
    try:
        with open(path, "rb") as fh:
            text = _decode(fh.read())
    except OSError as exc:
        raise ParseError(f"파일을 읽을 수 없다: {exc.strerror}") from exc

    builder = _Builder()
    # BOM을 지우지 않는다. 맨 앞 BOM 1개를 세지 않는 것은 hidden.py의 판정 규칙이고,
    # 여기서 지워버리면 뒤쪽에 숨겨 심은 BOM과 구분할 방법이 사라진다.
    for line in text.splitlines(keepends=True):
        body = line.rstrip("\r\n")
        ending = line[len(body):]
        builder.add_span(body, page=1)          # 줄 하나가 span 하나 (밀도 판정의 분모)
        builder.add_gap(ending, page=1)
    return builder.build("text", path, file_type)


# ---------------------------------------------------------------------------
# 이미지 — 텍스트가 없다. ml/의 CNN 파이프라인이 path를 받아간다.
# ---------------------------------------------------------------------------


def _load_image(path: str) -> ParsedDoc:
    return ParsedDoc(
        kind="image", raw_text="", spans=[], page_map=[], path=path, file_type="image"
    )


# ---------------------------------------------------------------------------
# PDF — get_texttrace()를 쓴다. get_text()가 아니다.
# ---------------------------------------------------------------------------
#
# get_text()로는 렌더 모드와 투명도가 안 나온다. 그 둘은 색 검사도 폰트 크기 검사도
# 통과하면서 화면에는 안 보이는 자리라, 숨은 명령을 심기에 가장 좋은 곳이다.
# get_texttrace()는 span마다 type(렌더모드)·opacity·color·size·bbox를 전부 준다.

# 같은 줄로 볼 세로 오차. 폰트 크기의 절반까지는 같은 줄로 친다.
_PDF_LINE_TOLERANCE_RATIO = 0.5
# 앞 조각과의 가로 간격이 공백 너비의 이 비율을 넘으면 공백을 하나 넣는다.
_PDF_SPACE_GAP_RATIO = 0.5


@dataclass
class _PdfItem:
    """get_texttrace() span 하나를 다루기 쉽게 정리한 것."""

    text: str
    bbox: tuple
    origin_y: float
    size: float
    color: str
    opacity: float
    render_mode: int
    spacewidth: float
    seqno: int
    char_x: list[float] | None


def _safe_chr(code) -> str:
    try:
        return chr(code)
    except (ValueError, TypeError):
        return "�"


def _char_x_edges(chars, direction) -> list[float] | None:
    """글자별 가로 경계 목록. 가로쓰기가 아니거나 좌표가 이상하면 None.

    None이면 locate.py가 span 사각형 전체로 물러선다 — 값보다 넓게 지우는 쪽이라
    개인정보가 남지는 않지만 옆 글자가 같이 지워질 수 있다.
    """
    if tuple(direction or (1.0, 0.0)) != (1.0, 0.0):    # 세로쓰기·회전된 텍스트
        return None
    edges: list[float] = []
    for char in chars:
        bbox = char[3]
        if not bbox or len(bbox) < 4:
            return None
        edges.append(float(bbox[0]))
    if not edges:
        return None
    edges.append(float(chars[-1][3][2]))
    # 왼쪽에서 오른쪽으로 정렬돼 있어야 구간을 잘라 쓸 수 있다.
    if any(b < a for a, b in zip(edges, edges[1:])):
        return None
    return edges


def _pdf_page_items(page) -> list[_PdfItem]:
    items: list[_PdfItem] = []
    for span in page.get_texttrace():
        chars = span.get("chars") or ()
        if not chars:
            continue
        text = "".join(_safe_chr(char[0]) for char in chars)
        if not text.strip():
            continue
        bbox = tuple(span.get("bbox") or (0.0, 0.0, 0.0, 0.0))
        items.append(
            _PdfItem(
                text=text,
                bbox=bbox,
                origin_y=float(chars[0][2][1]),
                size=float(span.get("size") or DEFAULT_FONT_SIZE),
                color=_pdf_color_to_hex(span.get("color")),
                opacity=float(span.get("opacity", 1.0)),
                render_mode=int(span.get("type", 0)),
                spacewidth=float(span.get("spacewidth") or 0.0),
                seqno=int(span.get("seqno", -1)),
                char_x=_char_x_edges(chars, span.get("dir")),
            )
        )
    return items


def _pdf_group_lines(items: list[_PdfItem]) -> list[list[_PdfItem]]:
    """읽는 순서로 줄을 묶는다.

    get_texttrace()는 파일에 쓰인 순서(seqno)로 준다. 그 순서가 곧 읽는 순서는
    아니라서(2단 편집, 나중에 끼워 넣은 문단) 좌표로 다시 정렬한다.
    """
    lines: list[list[_PdfItem]] = []
    for item in sorted(items, key=lambda it: (round(it.origin_y, 1), it.bbox[0])):
        tolerance = max(item.size * _PDF_LINE_TOLERANCE_RATIO, 2.0)
        if lines and abs(item.origin_y - lines[-1][0].origin_y) <= tolerance:
            lines[-1].append(item)
        else:
            lines.append([item])
    for line in lines:
        line.sort(key=lambda it: it.bbox[0])
    return lines


def _pdf_fills(page) -> list[tuple[int, object, str]]:
    """채움색이 있는 도형 목록. (그린 순서, 사각형, 색) — span의 배경색을 찾는 데 쓴다."""
    fills: list[tuple[int, object, str]] = []
    try:
        drawings = page.get_drawings()
    except Exception:      # 도형을 못 읽어도 텍스트 추출은 계속돼야 한다
        return fills
    for drawing in drawings:
        fill = drawing.get("fill")
        rect = drawing.get("rect")
        if fill is None or rect is None:
            continue
        opacity = drawing.get("fill_opacity")
        if opacity is not None and opacity <= 0.05:
            continue
        fills.append((int(drawing.get("seqno", -1)), rect, _pdf_color_to_hex(fill)))
    return fills


def _pdf_bg_at(fills, bbox: tuple, seqno: int) -> str | None:
    """글자 자리의 실제 배경 채움색. 없으면 None(= 모른다).

    흰색으로 가정하지 않는 이유: 파란 표 셀에 파란 글씨를 넣으면 "흰 배경 + 흰 글씨"만
    보는 검사를 그냥 통과한다. 글자보다 **먼저** 그려진 도형만 배경으로 친다
    (나중에 그려진 것은 글자를 덮은 것이지 배경이 아니다).
    """
    center_x = (bbox[0] + bbox[2]) / 2
    center_y = (bbox[1] + bbox[3]) / 2
    best: tuple[int, str] | None = None
    for fill_seqno, rect, color in fills:
        if fill_seqno > seqno:
            continue
        if not (rect.x0 <= center_x <= rect.x1 and rect.y0 <= center_y <= rect.y1):
            continue
        if best is None or fill_seqno > best[0]:
            best = (fill_seqno, color)
    return best[1] if best else None


def _load_pdf(path: str) -> ParsedDoc:
    import pymupdf

    try:
        document = pymupdf.open(path)
    except Exception as exc:
        raise ParseError(f"PDF를 열 수 없다: {type(exc).__name__}") from exc

    builder = _Builder()
    with document:
        if document.needs_pass:
            raise ParseError("암호가 걸린 PDF다")
        for page_number, page in enumerate(document, start=1):
            fills = _pdf_fills(page)
            for line in _pdf_group_lines(_pdf_page_items(page)):
                builder.newline(page_number)
                previous: _PdfItem | None = None
                for item in line:
                    if previous is not None:
                        gap = item.bbox[0] - previous.bbox[2]
                        if gap > previous.spacewidth * _PDF_SPACE_GAP_RATIO:
                            builder.add_gap(" ", page_number)
                    background = _pdf_bg_at(fills, item.bbox, item.seqno)
                    builder.add_span(
                        item.text,
                        page=page_number,
                        font_size=item.size,
                        color=item.color,
                        bg_color=background or DEFAULT_BG_COLOR,
                        bg_known=background is not None,
                        render_mode=item.render_mode,
                        opacity=item.opacity,
                        bbox=item.bbox,
                        char_x=item.char_x,
                    )
                    previous = item

    doc = builder.build("text", path, "pdf")
    # 텍스트 레이어가 없는 PDF = 스캔본이다. 이미지 파이프라인으로 보낸다.
    if len(doc.raw_text.strip()) < _PDF_SCANNED_TEXT_THRESHOLD:
        doc.kind = "image"
    return doc


# ---------------------------------------------------------------------------
# DOCX — 본문만 읽으면 안 된다. 표·머리말/꼬리말·텍스트상자·주석까지.
# ---------------------------------------------------------------------------
#
# python-docx의 Document.paragraphs는 본문 문단만 준다. 텍스트상자와 주석은 거기
# 없고, 숨은 명령은 바로 그런 자리에 심긴다. 그래서 문단 API 대신 각 파트의 XML을
# 문서 순서대로 훑는다.

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _w(tag: str) -> str:
    return _W + tag


# 파트 이름 -> (어디서 나온 텍스트인지, 읽는 순서). 본문을 먼저, 나머지를 뒤에 붙인다.
_DOCX_PART_RULES: tuple[tuple[re.Pattern, str, int], ...] = (
    (re.compile(r"^document\d*\.xml$"), "body", 0),
    (re.compile(r"^header\d*\.xml$"), "header", 1),
    (re.compile(r"^footer\d*\.xml$"), "footer", 2),
    (re.compile(r"^footnotes\.xml$"), "footnote", 3),
    (re.compile(r"^endnotes\.xml$"), "endnote", 3),
    (re.compile(r"^comments\.xml$"), "comment", 4),
)

_DOCX_TEXT_TAGS = (_w("t"), _w("delText"))
_DOCX_BREAK_TAGS = (_w("br"), _w("cr"))
_DOCX_PROPERTY_TAGS = (_w("rPr"), _w("pPr"), _w("tcPr"))

# w:highlight의 이름 있는 색 — 형광펜은 배경색으로 취급한다.
_DOCX_HIGHLIGHTS = {
    "white": "#ffffff", "black": "#000000", "yellow": "#ffff00", "green": "#00ff00",
    "cyan": "#00ffff", "magenta": "#ff00ff", "blue": "#0000ff", "red": "#ff0000",
    "darkBlue": "#000080", "darkCyan": "#008080", "darkGreen": "#008000",
    "darkMagenta": "#800080", "darkRed": "#800000", "darkYellow": "#808000",
    "darkGray": "#808080", "lightGray": "#c0c0c0",
}


def _xml_on(element) -> bool:
    """<w:vanish/>처럼 값이 없으면 켜진 것, w:val이 0/false면 꺼진 것."""
    if element is None:
        return False
    value = element.get(_w("val"))
    return value is None or value.lower() not in ("0", "false", "off")


def _ancestor(node, tag: str):
    current = node.getparent()
    while current is not None:
        if current.tag == tag:
            return current
        current = current.getparent()
    return None


def _docx_bg_color(node) -> str | None:
    """글자 자리의 배경 채움색. run 음영 -> 문단 음영 -> 표 셀 음영 순으로 올라간다."""
    current = node
    while current is not None:
        for property_tag in _DOCX_PROPERTY_TAGS:
            properties = current.find(property_tag)
            if properties is None:
                continue
            shading = properties.find(_w("shd"))
            if shading is None:
                continue
            fill = shading.get(_w("fill"))
            if fill and fill.lower() not in ("auto", "nil"):
                color = _hex_from_argb(fill)
                if color:
                    return color
        current = current.getparent()
    return None


def _docx_run_format(run, default_size: float) -> dict:
    """run 하나의 서식. 없는 값은 "보이는 쪽"의 기본값으로 채운다."""
    fmt = {
        "font_size": default_size,
        "color": DEFAULT_COLOR,
        "hidden_attr": False,
        "hidden_reason": "",
    }
    if run is None:
        return fmt
    properties = run.find(_w("rPr"))
    if properties is None:
        return fmt

    # w:vanish = 워드의 "숨김" 속성. webHidden은 웹 보기에서만 숨는 별개 속성이다.
    reasons = []
    if _xml_on(properties.find(_w("vanish"))):
        reasons.append("vanish")
    if _xml_on(properties.find(_w("webHidden"))):
        reasons.append("web_hidden")
    if reasons:
        fmt["hidden_attr"] = True
        fmt["hidden_reason"] = "+".join(reasons)

    size = properties.find(_w("sz"))
    if size is not None:
        try:
            fmt["font_size"] = float(size.get(_w("val"))) / 2   # w:sz는 half-point 단위다
        except (TypeError, ValueError):
            pass

    color = properties.find(_w("color"))
    if color is not None:
        parsed = _hex_from_argb(color.get(_w("val")))
        if parsed:
            fmt["color"] = parsed
    return fmt


def _docx_default_font_size(document) -> float:
    """styles.xml의 docDefaults에 적힌 기본 글자 크기. 없으면 11pt."""
    try:
        styles = document.styles.element
    except Exception:
        return DEFAULT_FONT_SIZE
    size = styles.find(f"{_w('docDefaults')}/{_w('rPrDefault')}/{_w('rPr')}/{_w('sz')}")
    if size is None:
        return DEFAULT_FONT_SIZE
    try:
        return float(size.get(_w("val"))) / 2
    except (TypeError, ValueError):
        return DEFAULT_FONT_SIZE


def _docx_parts(document) -> list[tuple[object, str]]:
    """텍스트가 들어 있는 파트만 골라 읽는 순서대로 돌려준다."""
    collected: list[tuple[int, str, object, str]] = []
    for part in document.part.package.iter_parts():
        if not hasattr(part, "element"):
            continue
        name = str(part.partname).rsplit("/", 1)[-1]
        for pattern, where, order in _DOCX_PART_RULES:
            if pattern.match(name):
                collected.append((order, name, part.element, where))
                break
    collected.sort(key=lambda item: (item[0], item[1]))
    return [(element, where) for _, _, element, where in collected]


def _load_docx(path: str) -> ParsedDoc:
    import docx

    try:
        document = docx.Document(path)
    except Exception as exc:
        raise ParseError(f"DOCX를 열 수 없다: {type(exc).__name__}") from exc

    default_size = _docx_default_font_size(document)
    builder = _Builder()

    for element, part_where in _docx_parts(document):
        previous_paragraph = None
        for node in element.iter():
            if not isinstance(node.tag, str):     # XML 주석 노드
                continue

            if node.tag in _DOCX_BREAK_TAGS:
                builder.add_gap("\n", page=1)
                continue
            if node.tag == _w("tab"):
                builder.add_gap("\t", page=1)
                continue
            if node.tag not in _DOCX_TEXT_TAGS:
                continue

            text = node.text or ""
            if not text:
                continue

            # 문단이 바뀌면 줄바꿈. 같은 문단 안의 run들은 붙여 쓴다 — 워드가 맞춤법
            # 검사나 서식 때문에 한 단어를 run 여러 개로 쪼개 놓는 일이 흔해서,
            # run 사이에 공백을 넣으면 "010-1234-" + "5678"이 안 잡힌다.
            paragraph = _ancestor(node, _w("p"))
            if paragraph is not previous_paragraph:
                builder.newline(page=1)
                previous_paragraph = paragraph

            run = _ancestor(node, _w("r"))
            fmt = _docx_run_format(run, default_size)

            where = part_where
            if _ancestor(node, _w("txbxContent")) is not None:
                where = "textbox"
            if node.tag == _w("delText"):
                # 변경내용 추적으로 지워진 글자. 화면에는 안 보이지만 파일에는 남아 있다.
                where = "deleted"
                fmt["hidden_attr"] = True
                fmt["hidden_reason"] = "+".join(filter(None, [fmt["hidden_reason"], "deleted"]))

            background = _docx_bg_color(run if run is not None else node)
            if background is None and run is not None:
                properties = run.find(_w("rPr"))
                highlight = properties.find(_w("highlight")) if properties is not None else None
                if highlight is not None:
                    background = _DOCX_HIGHLIGHTS.get(highlight.get(_w("val")) or "")

            builder.add_span(
                text,
                page=1,          # DOCX는 렌더링하기 전에는 페이지를 알 수 없다
                bg_color=background or DEFAULT_BG_COLOR,
                bg_known=background is not None,
                where=where,
                **fmt,
            )

    return builder.build("text", path, "docx")


# ---------------------------------------------------------------------------
# XLSX — 숨긴 시트, 숨긴 행/열, 셀 주석까지
# ---------------------------------------------------------------------------
#
# page에는 시트 번호(1부터)를 넣는다. 엑셀에 페이지 개념이 없어서다.

# 사용자 지정 서식 ";;;" = 셀에 값이 있어도 화면에는 아무것도 안 나온다.
_XLSX_BLANK_FORMAT = ";;;"

# 시트 상태 -> where. hidden과 veryHidden을 구분해서 넘긴다.
#
# 둘은 무게가 다르다. hidden은 엑셀에서 마우스 오른쪽 > "숨기기 취소"로 누구나 되돌릴
# 수 있어서 정상 문서에도 흔하다(계산용 시트를 접어둔 것). veryHidden은 그 메뉴에도
# 안 나와서 VBA나 XML을 직접 건드려야 만들어진다 — 실수로 그렇게 되는 일이 없다.
# 같은 "숨긴 시트"로 뭉뚱그리면 hidden.py가 정상 문서를 신고하거나, 반대로
# veryHidden의 강한 근거를 못 쓴다.
_XLSX_SHEET_WHERE = {
    "visible": "cell",
    "hidden": "sheet_hidden",
    "veryHidden": "sheet_very_hidden",
}

# theme1.xml의 색 순서는 dk1, lt1, dk2, lt2, accent1~6, hlink, folHlink인데
# 엑셀이 쓰는 인덱스는 앞의 두 쌍이 뒤바뀌어 있다 (0=lt1 배경, 1=dk1 글자).
_THEME_INDEX_ORDER = (1, 0, 3, 2, 4, 5, 6, 7, 8, 9, 10, 11)


def _xlsx_theme_colors(path: str) -> dict[int, str]:
    """워크북 테마 색.

    엑셀 색 선택기 첫 줄에서 흰색을 고르면 rgb가 아니라 theme 0으로 저장된다.
    이걸 못 읽으면 "흰 글꼴로 숨기기"를 통째로 놓친다.
    """
    try:
        with zipfile.ZipFile(path) as archive:
            data = archive.read("xl/theme/theme1.xml")
    except (KeyError, OSError, zipfile.BadZipFile):
        return {}
    try:
        root = ET.fromstring(data)
    except ET.ParseError:
        return {}

    namespace = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    scheme = root.find(f".//{namespace}clrScheme")
    if scheme is None:
        return {}

    ordered: list[str] = []
    for entry in scheme:
        srgb = entry.find(f"{namespace}srgbClr")
        system = entry.find(f"{namespace}sysClr")
        value = None
        if srgb is not None:
            value = _hex_from_argb(srgb.get("val"))
        elif system is not None:
            value = _hex_from_argb(system.get("lastClr"))
        ordered.append(value or DEFAULT_COLOR)

    colors: dict[int, str] = {}
    for index, source in enumerate(_THEME_INDEX_ORDER):
        if source < len(ordered):
            colors[index] = ordered[source]
    return colors


def _xlsx_color(color_obj, theme: dict[int, str]) -> str | None:
    """openpyxl의 Color 객체를 #rrggbb로. 알아낼 수 없으면 None."""
    if color_obj is None:
        return None
    color_type = getattr(color_obj, "type", None)
    if color_type == "rgb":
        return _hex_from_argb(getattr(color_obj, "rgb", None))
    if color_type == "theme":
        try:
            base = theme.get(int(color_obj.theme))
        except (TypeError, ValueError):
            return None
        return _apply_tint(base, getattr(color_obj, "tint", 0.0)) if base else None
    if color_type == "indexed":
        from openpyxl.styles.colors import COLOR_INDEX

        try:
            index = int(color_obj.indexed)
        except (TypeError, ValueError):
            return None
        if 0 <= index < len(COLOR_INDEX):
            return _hex_from_argb(COLOR_INDEX[index])
    return None


def _xlsx_hidden_columns(sheet) -> set[str]:
    """숨긴 열 문자 집합. 열은 묶음(min~max)으로 저장되는 경우가 있어 펼쳐서 담는다."""
    from openpyxl.utils import get_column_letter

    hidden: set[str] = set()
    for letter, dimension in sheet.column_dimensions.items():
        if not dimension.hidden:
            continue
        start, end = dimension.min, dimension.max
        if start and end:
            hidden.update(get_column_letter(i) for i in range(start, end + 1))
        else:
            hidden.add(letter)
    return hidden


def _xlsx_has_formulas(path: str) -> bool:
    """수식이 하나라도 있는 파일인가. 워크북을 두 번 여는 비용을 피하려고 먼저 확인한다."""
    try:
        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                if not (name.startswith("xl/worksheets/") and name.endswith(".xml")):
                    continue
                data = archive.read(name)
                if b"<f>" in data or b"<f " in data:
                    return True
    except (OSError, zipfile.BadZipFile):
        return False
    return False


def _cell_text(value) -> str:
    """셀 값을 문자열로.

    정수로 떨어지는 실수에 소수점을 붙이지 않는다 — 숫자로 입력된 전화번호가
    "1012345678.0"이 되면 정규식이 못 잡는다.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _load_xlsx(path: str) -> ParsedDoc:
    from openpyxl import load_workbook

    try:
        # data_only=True는 수식의 계산 결과를 준다. 결과가 저장돼 있지 않은 파일
        # (openpyxl로 만든 파일 등)에서는 None이 나오므로, 그때만 수식 원문을
        # 다시 읽어 채운다(formula_book). 값을 통째로 놓치면 탐지도 못 한다.
        workbook = load_workbook(path, data_only=True)
    except Exception as exc:
        raise ParseError(f"XLSX를 열 수 없다: {type(exc).__name__}") from exc

    theme = _xlsx_theme_colors(path)
    has_formulas = _xlsx_has_formulas(path)
    builder = _Builder()
    formula_book = None

    try:
        for sheet_number, sheet in enumerate(workbook.worksheets, start=1):
            # veryHidden은 엑셀 UI에서 "숨기기 취소"조차 보이지 않는다.
            sheet_hidden = sheet.sheet_state != "visible"
            sheet_where = _XLSX_SHEET_WHERE.get(sheet.sheet_state, "sheet_hidden")
            hidden_rows = {index for index, dim in sheet.row_dimensions.items() if dim.hidden}
            hidden_columns = _xlsx_hidden_columns(sheet)

            for row in sheet.iter_rows():
                row_started = False
                for cell in row:
                    text = _cell_text(cell.value)
                    if not text and has_formulas:
                        if formula_book is None:
                            formula_book = load_workbook(path, data_only=False)
                        text = _cell_text(formula_book[sheet.title][cell.coordinate].value)
                    if not text:
                        continue

                    if row_started:
                        builder.add_gap("\t", page=sheet_number)
                    else:
                        builder.newline(page=sheet_number)
                        row_started = True

                    font_color = _xlsx_color(getattr(cell.font, "color", None), theme)
                    background = None
                    if getattr(cell.fill, "patternType", None) == "solid":
                        background = _xlsx_color(getattr(cell.fill, "fgColor", None), theme)

                    reasons = []
                    if sheet_hidden:
                        reasons.append(sheet_where)
                    if cell.row in hidden_rows:
                        reasons.append("row_hidden")
                    if cell.column_letter in hidden_columns:
                        reasons.append("col_hidden")
                    if (cell.number_format or "").strip() == _XLSX_BLANK_FORMAT:
                        reasons.append("blank_format")

                    builder.add_span(
                        text,
                        page=sheet_number,
                        font_size=float(getattr(cell.font, "size", None) or DEFAULT_FONT_SIZE),
                        color=font_color or DEFAULT_COLOR,
                        bg_color=background or DEFAULT_BG_COLOR,
                        bg_known=background is not None,
                        hidden_attr=bool(reasons),
                        hidden_reason="+".join(reasons),
                        where=sheet_where,
                    )

                    # 셀 주석. 마우스를 올려야 보이지만 정상 문서에도 흔하므로
                    # hidden_attr로 찍지 않는다 — 판정은 hidden.py가 한다.
                    comment = getattr(cell, "comment", None)
                    if comment is not None and comment.text:
                        builder.add_gap("\t", page=sheet_number)
                        builder.add_span(
                            comment.text,
                            page=sheet_number,
                            hidden_attr=sheet_hidden,
                            hidden_reason=sheet_where if sheet_hidden else "",
                            where="comment",
                        )
    finally:
        workbook.close()
        if formula_book is not None:
            formula_book.close()

    return builder.build("text", path, "xlsx")


# ---------------------------------------------------------------------------
# 손으로 확인할 때 쓰는 진입점
# ---------------------------------------------------------------------------
#
#   uv run python -m backend.scanner.parser.parse <파일경로>
#
# 원문은 절대 찍지 않는다 (로그에 원문을 남기지 않는다 — 팀 규칙 2).


def _summary(doc: ParsedDoc) -> str:
    hidden = sum(1 for s in doc.spans if s.hidden_attr)
    invisible = sum(1 for s in doc.spans if s.render_mode == 3 or s.opacity == 0.0)
    tiny = sum(1 for s in doc.spans if s.font_size < 2.0)
    return (
        f"kind={doc.kind} type={doc.file_type} chars={len(doc.raw_text)} "
        f"spans={len(doc.spans)} pages={len(set(doc.page_map))}\n"
        f"hidden_attr={hidden} render_mode3_or_transparent={invisible} under_2pt={tiny}"
    )


if __name__ == "__main__":
    import sys

    if len(sys.argv) != 2:
        print("usage: python -m backend.scanner.parser.parse <file>")
        raise SystemExit(2)
    print(_summary(load(sys.argv[1])))
